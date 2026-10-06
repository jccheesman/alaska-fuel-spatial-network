#!/usr/bin/env python3
"""Stage 06 — build the owner's review queue from the structural and boundary flags.

One row per (record, flag) that is NOT already covered by a correction (approved / pending /
rejected) or a remote-site exception, with evidence, a suggested fix where a detector has
one, a map link, and blank decision columns. Nothing here changes data: the owner's
answers go back through `record-facility-corrections` into corrections.csv / remote_sites.csv.

Writes outputs/00_inventory_qc/<release>/review_queue.csv (always) and review_queue.xlsx
(when openpyxl is installed: dropdown decisions, map hyperlinks, frozen header).

Run:  python workflows/00_inventory_qc/06_review_queue.py [--release 2025]
"""
from __future__ import annotations

import argparse
import sys

import numpy as np
import pandas as pd

from _qc import OUT, latest_release, read_csv, write_csv

DECISIONS = ["approve", "approve with my value", "reject", "ask publisher", "remote site (keep as labelled)", "skip for now"]
ORDER = {"high": 0, "review": 0, "medium": 1, "low": 2, "info": 3}


def build(label: str) -> pd.DataFrame:
    out = OUT / label
    fac = read_csv(out / "facilities_normalised.csv").set_index("record_id")
    S = read_csv(out / "flags_structural.csv")
    S = S[S["severity"] != "info"]
    S = S.assign(source="structural", priority=S["severity"],
                 suggestion=np.where(S["suggested_value"].notna(), S["suggested_field"].fillna("") + " -> " + S["suggested_value"].fillna(""), ""))
    B = read_csv(out / "flags_boundary.csv")
    B = B.assign(source="boundary", detector="boundary_" + B["outcome"], priority=B["priority"],
                 detail=("inside " + B["inside_boundary_of"].fillna("no boundary") + "; " + B["own_boundary_km"].astype(str) + " km from own boundary"),
                 group_key="", suggestion="")
    cols = ["record_id", "source", "detector", "priority", "detail", "suggestion", "group_key", "covered_by"]
    Q = pd.concat([S[cols], B[cols]], ignore_index=True)
    Q = Q[Q["covered_by"].isna() | (Q["covered_by"] == "")].drop(columns="covered_by")
    for c in ("community_name", "entity_name", "latitude", "longitude", "tank_farm_id", "total_capacity", "source_row"):
        Q[c] = Q["record_id"].map(fac[c]) if c in fac.columns else np.nan
    Q["map"] = "https://www.google.com/maps?q=" + Q["latitude"].astype(str) + "," + Q["longitude"].astype(str)
    Q.loc[Q["latitude"].isna(), "map"] = ""
    for c in ("decision", "your_value", "notes", "reviewed_by", "reviewed_on"):
        Q[c] = ""
    Q["_o"] = Q["priority"].map(ORDER).fillna(9)
    Q = Q.sort_values(["_o", "group_key", "record_id", "detector"]).drop(columns="_o").reset_index(drop=True)
    front = ["record_id", "source_row", "community_name", "entity_name", "latitude", "longitude", "tank_farm_id", "total_capacity",
             "source", "detector", "priority", "detail", "suggestion", "group_key", "map", "decision", "your_value", "notes", "reviewed_by", "reviewed_on"]
    return Q[front]


def write_xlsx(Q: pd.DataFrame, path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return False
    wb = Workbook(); ws = wb.active; ws.title = "Review queue"
    hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="305496")
    inp = PatternFill("solid", fgColor="FFFF00"); f = Font(name="Arial", size=10); link = Font(name="Arial", size=10, color="0563C1", underline="single")
    widths = {"detail": 80, "suggestion": 30, "entity_name": 30, "notes": 40, "community_name": 18, "detector": 20, "group_key": 14, "your_value": 22}
    for j, c in enumerate(Q.columns, 1):
        x = ws.cell(row=1, column=j, value=c); x.font = hdr; x.fill = fill; x.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 13)
    inputs = {"decision", "your_value", "notes", "reviewed_by", "reviewed_on"}
    for i, row in enumerate(Q.itertuples(index=False), 2):
        for j, (c, v) in enumerate(zip(Q.columns, row), 1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            if c == "map" and v:
                x = ws.cell(row=i, column=j, value=f'=HYPERLINK("{v}","map")'); x.font = link
            else:
                x = ws.cell(row=i, column=j, value=v); x.font = f
                x.alignment = Alignment(wrap_text=(c == "detail"), vertical="top")
            if c in inputs:
                x.fill = inp
    dcol = get_column_letter(list(Q.columns).index("decision") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dv.add(f"{dcol}2:{dcol}{len(Q)+1}"); ws.add_data_validation(dv)
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(Q.columns))}{len(Q)+1}"
    wb.save(path)
    return True


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--release", default=None)
    a = ap.parse_args(argv)
    label = a.release or latest_release()
    Q = build(label)
    write_csv(Q, OUT / label / "review_queue.csv")
    xl = write_xlsx(Q, OUT / label / "review_queue.xlsx")
    by = Q.groupby(["priority", "detector"]).size()
    print(f"release {label}: {len(Q)} open flags on {Q['record_id'].nunique()} records" + ("" if xl else " (openpyxl not installed: CSV only)"))
    for (p, d), n in by.items():
        print(f"  {p:7s} {d:22s} {n:4d}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
